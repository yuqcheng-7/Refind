from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('output/Refind拾藏PRD_V1.0.docx')
SOURCE_MD = Path('output/Refind拾藏PRD_V1.0.md')

BLUE = '1F4E79'
LIGHT_BLUE = 'EAF3F8'
GRAY = 'F3F5F7'
DARK = '1F2937'
MUTED = '5F6B76'

def set_font(run, name='Arial Unicode MS', size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_table_borders(table, color='D4DCE4'):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)
        borders.append(el)
    tblPr.append(borders)

def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'), str(int(width * 1440))); tcW.set(qn('w:type'), 'dxa')

def add_para(doc, text='', style=None, bold_lead=None, quote=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    if quote:
        p.paragraph_format.left_indent = Inches(.25)
        p.paragraph_format.right_indent = Inches(.15)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); set_font(r, size=10.5, bold=True, color=BLUE)
        r = p.add_run(text[len(bold_lead):]); set_font(r, size=10.5, color=DARK)
    else:
        r = p.add_run(text); set_font(r, size=10.5, color=DARK)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text); set_font(r, size=10.3, color=DARK)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text); set_font(r, size=10.3, color=DARK)
    return p

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    set_table_borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c, LIGHT_BLUE); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        r=p.add_run(h); set_font(r,size=9.5,bold=True,color=BLUE)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
            r=p.add_run(str(val)); set_font(r,size=9.2,color=DARK)
    if widths: set_table_widths(t,widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next=True
    r=p.add_run(text); set_font(r, size={1:16,2:13,3:11.5}[level], bold=True, color=BLUE if level<3 else DARK)
    return p

def add_code(doc, lines):
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(7)
    r=p.add_run('\n'.join(lines)); set_font(r,name='Menlo',size=9.2,color='374151')
    return p

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)

# styles
normal=doc.styles['Normal']; normal.font.name='Arial Unicode MS'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS'); normal.font.size=Pt(10.5)
for s, size in [('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
    st=doc.styles[s]; st.font.name='Arial Unicode MS'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(BLUE); st.font.bold=True
    st.paragraph_format.space_before=Pt(15 if s=='Heading 1' else 10); st.paragraph_format.space_after=Pt(6)

# header/footer
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('Refind · 拾藏 PRD（V1.0）'); set_font(r,size=8.5,color=MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Refind · AI 个人知识助手'); set_font(r,size=8.5,color=MUTED)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(54); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Refind · 拾藏'); set_font(r,size=28,bold=True,color=BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(7)
r=p.add_run('AI 个人知识助手 PRD'); set_font(r,size=17,bold=True,color=DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('V1.0｜个人知识库 · RAG 对话 · 来源溯源'); set_font(r,size=10.5,color=MUTED)
doc.add_paragraph().paragraph_format.space_after=Pt(12)
t=doc.add_table(rows=3,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; set_table_borders(t,'DCE3EA'); set_table_widths(t,[1.25,4.3])
for i,(a,b) in enumerate([('版本','V1.0'),('产品定位','面向工作、学习与研究场景的 AI 个人知识助手'),('文档日期','2026-08-13')]):
    shade(t.rows[i].cells[0],LIGHT_BLUE)
    for j,val in enumerate([a,b]):
        c=t.rows[i].cells[j]; set_cell_margins(c); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(val); set_font(r,size=10,bold=(j==0),color=BLUE if j==0 else DARK)
doc.add_page_break()

heading(doc,'一、项目概述',1)
heading(doc,'1.1 项目背景',2)
add_para(doc,'在工作、学习和研究过程中，用户会在小红书、微信公众号、知乎、浏览器网页等渠道收藏案例、文章、教程和观点。')
add_para(doc,'但当需要完成竞品分析、方案撰写、概念学习等具体任务时，用户往往不记得资料保存在哪里，也难以快速判断哪些历史收藏与当前问题相关，仍需花大量时间重新搜索、逐篇阅读和手动整理。')
add_para(doc,'因此，用户的核心痛点并非“无法收藏”，而是：',bold_lead='因此，')
add_para(doc,'已积累大量历史资料，却难以在需要时高效将其转化为解决具体任务的依据与思路。',quote=True)
heading(doc,'1.2 项目目标',2)
heading(doc,'1.2.1 业务目标',3)
add_para(doc,'面向高频收集外部信息，且需要在工作、学习或研究中持续调用资料的个人知识使用者，构建以“个人知识高效调用”为核心的 AI 产品。')
add_para(doc,'通过降低用户从资料收集、找回、理解到应用的时间成本，建立“随手沉淀资料、遇到问题优先调用个人知识库”的使用习惯，并为后续扩展账号同步、多模态资料导入、共享知识库等能力奠定基础。')
heading(doc,'1.2.2 产品目标',3)
add_para(doc,'Refind 因此将分散收藏沉淀为个人知识库，并通过 AI 帮助用户提升资料使用效率。用户保存资料后，AI 自动提炼摘要与标签，帮助资料形成基础结构；当用户在工作、学习或研究中遇到具体问题时，可直接用自然语言向个人知识库提问。AI 从相关资料中检索内容，归纳为回答、分析思路或信息线索，并展示每个关键结论的来源，方便用户核验和继续使用。')
add_para(doc,'用户不再只是“找回一篇收藏”，而是能够基于过去积累的资料，更快地理解问题、形成思路并推进当前任务。')
add_para(doc,'为实现上述目标，产品能力将围绕以下方向展开：')
for b in ['资料沉淀：将分散收藏转化为可管理的个人知识资产；','知识组织与找回：支持按主题管理资料，并快速定位相关内容；','AI 知识调用：支持用户围绕具体问题，与个人知识库进行对话；','可信回答与核验：确保 AI 回答有资料依据，用户可追溯、核验并继续使用。']: add_bullet(doc,b)
heading(doc,'1.2.3 产品定位',3)
add_para(doc,'Refind 是一款面向工作、学习与研究场景的 AI 个人知识助手。它帮助用户将日常收藏的网页与社媒资料沉淀为可管理、可检索、可对话的个人知识库；当用户遇到具体问题时，AI 仅基于个人资料提供可溯源的回答与思路。',quote=True)
heading(doc,'1.3 V1 范围',2)
add_table(doc,['本期实现','本期不实现'],[[
'多知识库创建、编辑、删除与资料管理\n默认知识库\n网页/社媒链接保存\n链接内容解析及 AI 自动结构化\n全局与知识库内语义搜索\n全局 AI 对话与 #标签 聚焦\n来源引用、原文片段查看与资料不足提示\n一键生成脑图',
'登录、注册、账号隔离与跨端同步\n联网搜索及知识库外信息补充\n团队协作、共享知识库与权限体系\n文件、图片、音频等多模态资料导入\n复杂 Agent 工作流\n知识图谱、自动主题聚类与内容推荐']], [3.25,3.25])

heading(doc,'二、用户与场景分析',1)
heading(doc,'2.1 目标用户',2)
add_para(doc,'Refind 面向高频收集外部信息，并需要将资料用于实际工作、学习或研究任务的个人知识使用者。V1 优先覆盖以下两类用户：')
add_table(doc,['用户类型','典型人群','资料类型与调用任务'],[
['职场知识工作者','产品、运营、市场、咨询、设计等从业者','收藏行业文章、竞品案例、用户洞察、活动玩法与方法论；用于方案撰写、竞品分析、项目调研、内容策划或汇报准备。'],
['研究与学习人群','学生、研究生、备考者及转行者','收藏课程资料、论文、教程、岗位信息、面试经验和学习方法；用于理解概念、梳理主题、准备考试、开展研究或解决学习问题。']], [1.25,1.75,3.5])
add_para(doc,'两类用户的共同需求是：',bold_lead='两类用户')
add_para(doc,'当我面对一个具体问题时，我希望直接向自己积累的资料提问，获得有依据、可核验的回答与思路，而不是重新搜索和逐篇翻找。',quote=True)
heading(doc,'2.2 用户痛点',2)
add_table(doc,['用户阶段','用户行为','核心痛点'],[
['信息发现','在社媒、网页或文章中发现有价值内容','资料来源分散，难以统一沉淀。'],
['信息保存','使用平台收藏、复制链接或截图保存','缺少统一结构，后续难以回忆内容价值与适用场景。'],
['资料管理','历史收藏不断累积','主题混杂，难以按任务或领域组织。'],
['资料找回','围绕当前任务寻找过往资料','不记得标题、平台或关键词，翻找成本高。'],
['知识理解','阅读多篇资料并判断其价值','需要手动通读、提炼、比较和归纳。'],
['任务应用','将资料用于方案、调研或学习','难以快速形成有依据的答案与思路；通用 AI 回答又难以核验。']], [1.05,1.65,3.8])
heading(doc,'2.3 核心使用场景',2)
for title,problem,q,reply in [
('场景一：工作调研与方案准备','用户需要完成竞品分析、活动方案、内容策划或项目调研，但无法快速找回过去收藏的案例与方法论。','根据我收藏的资料，会员活动可以从哪些维度设计？','Refind 从个人知识库中检索相关案例与方法论，输出可参考的分析维度、案例做法与思考方向，并附带对应资料来源。'),
('场景二：学习与概念理解','用户学习某一主题时，已经收藏了教程、文章或论文，但需要理解概念关系或解决具体问题。','根据我收藏的 RAG 资料，召回和重排序分别解决什么问题？','Refind 基于相关资料进行归纳，说明概念关系与适用场景，并展示支撑回答的原文片段。'),
('场景三：资料回顾与定位','用户只记得过去保存过某类资料，但无法记起准确标题或关键词。','提高新用户留存的方法','Refind 返回语义相关的资料，并展示标题、匹配片段、标签和所属知识库，帮助用户快速定位需要的内容。'),
]:
    heading(doc,title,3); add_para(doc,problem); add_para(doc,q,quote=True); add_para(doc,reply)
heading(doc,'2.4 核心用户流程',2)
add_code(doc,['发现资料 → 保存至知识库 → 自动结构化 → 在任务中搜索或提问 → 获得有来源的资料/回答 → 核验并继续使用'])

heading(doc,'三、产品整体架构',1)
heading(doc,'3.1 产品功能架构',2)
add_code(doc,['Refind · AI 个人知识助手','├── 资料沉淀：链接保存｜链接解析｜AI 摘要与标签','├── 知识库管理：知识库列表｜默认知识库｜资料移动、编辑、删除','├── 知识组织与找回：全局语义搜索｜知识库内语义搜索｜搜索建议浮层','└── AI 知识调用：全局对话｜#标签聚焦｜RAG 检索｜来源引用｜资料不足提示'])
heading(doc,'3.2 模块说明',2)
add_table(doc,['一级模块','模块目标','核心功能'],[
['资料沉淀','让用户低成本把外部资料带入 Refind','粘贴链接、解析标题与内容、生成摘要和标签、选择知识库保存。'],
['知识库管理','让用户按主题组织、维护历史资料','创建知识库、默认知识库、资料查看、移动、编辑、删除。'],
['知识组织与找回','帮助用户快速定位某篇相关资料','全局及库内语义搜索、搜索建议浮层、匹配片段展示。'],
['AI 知识调用','帮助用户基于已有资料理解问题、形成思路','全局对话、标签聚焦、检索回答、来源溯源、资料不足提示。']], [1.25,1.75,3.5])
heading(doc,'3.3 信息架构',2)
add_code(doc,['首页：全局知识库','├── 全局 AI 对话入口','├── 全局语义搜索','├── 全局添加资料','└── 知识库列表','    ├── 创建新知识库','    ├── 默认知识库','    └── 单个知识库详情：资料列表｜库内搜索｜添加资料｜资料详情'])

heading(doc,'四、详细功能需求',1)
heading(doc,'4.1 首页：知识库总览',2)
add_para(doc,'页面目标：提供知识库浏览、全局资料搜索及 AI 对话入口。页面采用左侧知识库列表、右侧当前知识库资料卡片区的布局；顶部提供全局搜索框与添加资料入口，AI 对话入口位于页面显著位置。')
add_table(doc,['区域','详细需求'],[
['左侧知识库列表','展示全部知识库名称、资料数量及当前选中状态；支持点击切换知识库；提供创建知识库按钮；默认知识库固定展示，可重命名但不可删除。'],
['右侧资料区域','展示当前知识库名称、简介、资料数量；提供添加资料与排序入口；以卡片展示标题、来源平台、摘要、标签、保存时间和正文节选预览；点击进入资料详情。'],
['全局搜索','支持语义搜索全部知识库资料；输入后展示搜索建议浮层；搜索结果展示所属知识库，点击跳转资料详情。'],
['空状态','首次无知识库时展示“创建第一个知识库，开始沉淀你的资料”；主按钮为创建知识库，次入口为添加资料。']], [1.25,5.25])
heading(doc,'4.2 创建知识库弹窗',2)
add_table(doc,['字段/功能','规则'],[['知识库名称','必填，1–30 字。'],['知识库简介','选填，最长 100 字。'],['创建按钮','名称合法时可点击。'],['取消按钮','关闭弹窗，不保存输入内容。']], [1.45,5.05])
add_para(doc,'创建成功后，左侧知识库列表新增并高亮当前知识库；右侧切换为该知识库的空资料列表，并展示“添加第一篇资料”引导。')
heading(doc,'4.3 添加资料页',2)
heading(doc,'4.3.1 链接输入状态',3)
add_table(doc,['功能','规则'],[['链接输入框','必填，仅支持 http/https 链接。'],['所属知识库','从知识库内进入时默认选中当前知识库；从首页进入时可选择。'],['开始解析','链接合法后可点击。'],['取消','返回上一页面，不保存资料。']], [1.45,5.05])
add_para(doc,'未选择所属知识库时，资料保存至默认知识库；默认知识库不存在时自动创建。')
heading(doc,'4.3.2 解析中状态',3)
add_para(doc,'用户点击开始解析后，展示不可重复提交的加载状态，并按以下顺序显示当前处理步骤：')
add_code(doc,['正在解析链接内容 → 正在识别标题与来源平台 → 正在提取正文或内容节选 → 正在生成摘要与标签 → 解析完成'])
heading(doc,'4.3.3 资料预览状态',3)
add_para(doc,'解析完成后展示原始链接、自动识别的来源平台与标题、正文或内容节选、AI 生成的摘要与标签、所属知识库。用户可编辑标题、摘要、标签和所属知识库；来源平台默认不可编辑。提供保存资料、重新解析与取消入口。')
heading(doc,'4.3.4 解析失败状态',3)
add_para(doc,'当 URL 合法但无法完成内容解析时，展示失败说明、重新解析和仅保存链接入口。选择仅保存链接后：标题默认使用链接域名，来源平台标记为“其他”，正文节选、摘要和标签为空，资料状态标记为“待补充内容”；该资料后续支持重新解析，在解析成功前不参与搜索和 AI 对话。')
heading(doc,'4.4 知识库资料列表',2)
add_para(doc,'资料卡片展示标题、来源平台、摘要、标签、保存时间、正文节选预览及待补充内容状态（如适用）。默认按最近添加时间倒序展示。')
add_table(doc,['排序方式','规则'],[['最近添加','从新到旧、从旧到新。'],['标题','A–Z、Z–A。'],['来源平台','按平台名称排序。']], [1.45,5.05])
add_para(doc,'点击排序入口打开排序浮层，选择后即时刷新列表。知识库内搜索仅检索当前知识库；未主动选择排序时按相关度展示结果，选择排序后按所选规则排列。待补充内容资料不参与搜索。')
heading(doc,'4.5 搜索建议浮层',2)
add_table(doc,['状态','页面表现'],[['输入为空','不展示浮层。'],['搜索中','展示加载状态。'],['有结果','展示资料列表。'],['无结果','展示“未找到相关资料”。'],['关闭','点击空白处或关闭按钮关闭浮层。']], [1.45,5.05])
add_para(doc,'每条结果展示资料标题、所属知识库、匹配内容片段、标签和来源平台；点击条目跳转资料详情页。')
heading(doc,'4.6 资料详情页',2)
add_para(doc,'展示标题、所属知识库、来源平台、保存时间、原始链接、AI 摘要、标签、正文或内容节选及资料状态。支持查看原文、编辑资料、移动资料、重新解析和删除资料；删除需二次确认，并删除关联片段与向量索引。待补充内容资料需突出展示重新解析入口。')
heading(doc,'4.7 AI 对话页',2)
add_para(doc,'用户可输入自然语言问题，并通过一个或多个 #标签 聚焦资料范围。发送后展示检索与生成状态。回答区展示回答、分析思路或信息线索，顶部显示“本次参考 X 个知识库、Y 篇资料”，关键结论附引用标记，并提供有帮助/无帮助反馈；当资料充足且回答存在层级结构时展示生成脑图入口。')
add_para(doc,'资料不足时，展示已检索到的相关内容、当前缺少的信息及建议补充的资料类型；不输出无依据的完整结论，也不展示生成脑图入口。')
heading(doc,'4.8 来源引用卡片',2)
add_para(doc,'用户点击回答中的引用标记后，展示来源卡片：资料标题、所属知识库、来源平台、支持该结论的原文片段、查看原文入口及关闭按钮。点击空白处、关闭按钮或切换其他引用时关闭或切换卡片。')
heading(doc,'4.9 一键生成脑图',2)
add_para(doc,'入口位于符合条件的 AI 回答底部。脑图仅基于本轮回答及其已引用资料生成，以问题或回答主题为中心节点，向下展示核心结论、子观点、案例或行动建议。节点支持展开与收起；有来源依据的节点展示引用标记，点击可查看对应来源。生成失败时提示“脑图生成失败，请重试”。')

heading(doc,'五、核心用户流程',1)
heading(doc,'5.1 首次使用：创建知识库并添加第一篇资料',2)
add_code(doc,['进入首页空状态 → 创建知识库 → 左侧展示并高亮当前知识库 → 右侧展示空资料列表','→ 引导添加第一篇资料 → 粘贴链接并开始解析 → 查看资料预览 → 确认保存 → 返回知识库，展示资料卡片'])
heading(doc,'5.2 创建知识库',2)
add_code(doc,['点击“创建知识库” → 输入名称与简介（选填） → 点击创建 → 创建成功并进入该知识库','→ 左侧列表新增并高亮当前知识库 → 右侧展示空资料列表与“添加第一篇资料”引导'])
heading(doc,'5.3 新建资料',2)
add_code(doc,['从首页或知识库内点击“添加资料” → 粘贴链接 → 选择归属知识库（可选） → 开始解析','→ 展示解析进度 → 解析完成，进入资料预览页 → 用户确认或编辑标题、摘要、标签、归属','→ 保存资料 → 返回所属知识库，新增资料卡片'])
heading(doc,'5.4 浏览、搜索与定位资料',2)
add_code(doc,['进入首页或知识库详情 → 在搜索框输入自然语言 → 展示搜索建议浮层 → 点击命中资料 → 进入资料详情页或查看原文'])
heading(doc,'5.5 全局 AI 对话',2)
add_code(doc,['在首页输入具体问题 → 可选输入 #标签 聚焦资料范围 → 系统检索全部知识库或指定标签下资料','→ AI 基于资料生成回答与思路 → 展示参考知识库、资料与引用 → 用户核验来源或继续追问'])

heading(doc,'六、关键业务规则',1)
heading(doc,'6.1 首页与知识库空状态',2)
for b in ['首次进入且尚未创建任何知识库时，首页展示空状态；主文案为“创建第一个知识库，开始沉淀你的资料”。','主按钮为“创建知识库”，次入口为“添加资料”；从首页添加资料时若默认知识库不存在，系统自动创建默认知识库作为归属。','创建知识库后，左侧展示知识库列表并高亮当前知识库；右侧展示资料卡片区域。当前知识库无资料时，显示“知识库还是空的，添加第一篇资料开始沉淀知识”及添加资料按钮。']: add_bullet(doc,b)
heading(doc,'6.2 知识库创建与管理',2)
for b in ['知识库名称必填，1–30 字；简介选填，最长 100 字；允许创建同名知识库。','列表展示名称、资料数量和最近更新时间；支持编辑知识库名称与简介。','删除知识库需二次确认；库内资料自动移动至默认知识库，避免资料丢失；默认知识库可重命名但不可删除。']: add_bullet(doc,b)
heading(doc,'6.3 资料归属规则',2)
for b in ['从知识库内添加资料时默认归属当前知识库；从首页添加时用户可选择归属。','未选择知识库时自动保存至默认知识库；默认知识库不存在时自动创建。','已保存资料支持移动至其他知识库。']: add_bullet(doc,b)
heading(doc,'6.4 链接解析、预览与保存',2)
for b in ['系统自动识别来源平台，抓取或识别页面标题，提取正文或内容节选，并生成 AI 摘要与标签。','解析完成后进入预览页；用户确认后资料才正式保存。用户可编辑标题、摘要、标签和所属知识库，来源平台不可编辑。','标题为必填项；摘要、标签和正文节选允许为空。保存中按钮展示加载并禁止重复提交；成功后提示“资料已保存”并返回所属知识库。','保存失败时停留预览页并保留全部编辑内容，提示“保存失败，请重试”；网络异常提示检查网络，服务异常提示稍后重试。','同一知识库中保存重复链接时提示“该资料已存在”，并提供查看已有资料入口。']: add_bullet(doc,b)
heading(doc,'6.5 解析失败与仅保存链接',2)
for b in ['URL 合法但标题、来源或正文解析失败时，用户可选择重新解析或仅保存链接。','仅保存链接时，保存 URL、资料归属与保存时间；标题使用链接域名，来源平台标记为“其他”，正文节选、摘要和标签为空，状态标记为“待补充内容”。','待补充内容资料支持后续重新解析；成功前不参与搜索与 AI 对话。若资料写入失败，系统不得提示保存成功，应保留页面信息供用户重试或复制链接。']: add_bullet(doc,b)
heading(doc,'6.6 资料列表与排序',2)
for b in ['资料以卡片展示标题、来源平台、摘要、标签、保存时间与正文节选预览，默认按最近添加时间倒序。','支持最近添加（新到旧/旧到新）、标题（A–Z/Z–A）和来源平台排序；选择排序后列表即时刷新并展示当前排序。','排序可与知识库内语义搜索同时生效：未指定排序时按相关度展示，指定排序后按所选规则排列。']: add_bullet(doc,b)
heading(doc,'6.7 搜索建议浮层',2)
for b in ['支持全局语义搜索和知识库内语义搜索；全局搜索检索全部知识库，库内搜索仅检索当前知识库。','输入自然语言后展示搜索建议浮层；每条展示标题、所属知识库、匹配片段、标签和来源平台；点击跳转资料详情。','输入为空不展示；加载中展示加载状态；无命中展示“未找到相关资料”；点击空白处或关闭按钮关闭。待补充内容资料不参与搜索。']: add_bullet(doc,b)
heading(doc,'6.8 #标签聚焦提问',2)
for b in ['#标签仅匹配资料标签，不匹配知识库名称；用户可添加一个或多个标签。','多个标签为 AND 关系，仅检索同时具备全部标签的资料；未添加时默认检索全部知识库。','标签不存在或范围内无相关资料时提示当前范围无相关资料；资料不足时不得自动扩大检索范围。']: add_bullet(doc,b)
heading(doc,'6.9 AI 对话与来源溯源',2)
for b in ['AI 对话从首页全局入口发起，只基于知识库可用资料回答，不联网、不使用知识库外信息。','每轮提问均重新检索相关资料；回答顶部展示参考知识库数量与资料数量。','关键结论、观点和案例须附来源引用；点击引用展示资料标题、所属知识库、支持结论的原文片段及查看原文入口。','AI 可归纳多篇资料，但不得将推断表述为既有事实；资料明显冲突时需说明差异并分别标注来源。','用户可继续追问，后续回答仍基于知识库资料与当前对话上下文。']: add_bullet(doc,b)
heading(doc,'6.10 资料不足处理',2)
add_para(doc,'以下情况触发资料不足：未检索到相关资料、资料相关性较弱、资料仅覆盖部分问题，或存在无法判断的冲突。此时 AI 不输出看似完整但无依据的结论，应说明已找到的内容、缺少的信息及建议补充的资料类型。')

heading(doc,'七、AI 能力方案与验收标准',1)
heading(doc,'7.1 AI 技术方案',2)
add_para(doc,'Refind V1 采用 RAG（检索增强生成）架构。AI 不直接基于通用知识回答，而是先从知识库中检索相关资料，再基于检索结果生成回答，并为关键结论提供来源引用。')
add_code(doc,['用户保存链接 → 解析网页内容 → DeepSeek 生成摘要与标签 → 正文切分为资料片段 → Embedding 向量化 → pgvector 建立索引','用户提问 → 解析 #标签 → 问题向量化 → 标签过滤 → pgvector 相似度检索 → DeepSeek 基于片段回答 → 绑定来源引用'])
heading(doc,'7.2 模型与技术组件分工',2)
add_table(doc,['组件','技术选型','主要职责'],[
['生成模型','DeepSeek Chat','生成资料摘要、标签，以及基于检索片段生成最终回答。'],
['向量模型','阿里云百炼 text-embedding-v4','将资料片段与用户问题转换为向量，用于中文语义检索。'],
['向量存储与检索','Supabase Postgres + pgvector','存储向量，按相似度检索相关资料片段，并结合标签过滤。'],
['关系数据存储','Supabase Postgres','保存知识库、资料、标签、资料片段、引用关系与对话记录。'],
['后端编排','Supabase Edge Functions','串联链接解析、结构化、向量化、检索、回答生成与引用返回。']], [1.25,2.15,3.1])
heading(doc,'7.3 资料入库与索引构建',2)
for n in ['解析链接，获取标题、来源平台、正文或内容节选。','调用 DeepSeek Chat 生成资料摘要与标签。','对正文清洗并按段落或固定长度切分为多个资料片段。','调用 text-embedding-v4 为每个资料片段生成向量。','将资料片段、向量及资料 ID、知识库 ID、标题、平台、URL、标签和原文位置等元数据保存至 pgvector。']: add_number(doc,n)
add_para(doc,'资料编辑、重新解析或删除后，需同步更新或删除相关资料片段与向量，避免过期内容参与检索。待补充内容资料不生成向量。')
heading(doc,'7.4 提问理解与检索规则',2)
add_para(doc,'系统将用户问题转换为向量，在可用资料片段中召回语义最相关内容。含 #标签时，先按资料标签过滤，再执行向量相似度检索；多个标签采用 AND 关系，范围内资料不足时不扩大范围。不含 #标签时，在全部知识库可用资料片段中检索，并保留所属知识库信息用于回答展示。')
heading(doc,'7.5 回答生成规则',2)
for b in ['仅使用检索到的资料片段与当前对话上下文；不联网，不使用知识库外信息补充。','优先直接回应用户问题，可归纳多篇资料中的共同观点、案例或差异。','对未直接说明的推断，使用“根据现有资料可推断”“可进一步考虑”等表达。','对资料间明显冲突说明差异，不强行给唯一结论。','每个关键事实、案例、观点或结论必须关联至少一条来源；不得将资料中不存在的内容描述为既有事实。']: add_bullet(doc,b)
heading(doc,'7.6 来源引用与溯源',2)
for b in ['关键结论后展示引用标记，如 [1]、[2]；同一结论可关联多条来源。','点击引用打开来源卡片，展示资料标题、所属知识库、来源平台、支持结论的原文片段和查看原文入口。','回答顶部展示“本次参考 X 个知识库、Y 篇资料”，用户可展开查看本轮使用的全部资料。']: add_bullet(doc,b)
heading(doc,'7.7 资料不足判断',2)
add_para(doc,'未检索到相关资料、召回资料相关性低、资料只覆盖问题的一部分、资料内容过短或存在无法判断的冲突时，AI 应优先说明边界。回复需包含已找到的相关内容、现有资料无法覆盖的部分及建议补充的资料类型。')
heading(doc,'7.8 AI 验收标准',2)
add_table(doc,['能力','验收标准'],[
['摘要生成','能概括资料核心信息；用户可编辑；生成失败不阻断资料保存。'],
['标签生成','标签反映资料主题；可用于 #标签 过滤；用户可编辑。'],
['语义检索','输入自然语言后，可召回标题不完全匹配但主题相关的资料。'],
['标签聚焦','使用多个 #标签 时，仅检索同时具备全部标签的资料。'],
['知识库问答','回答仅使用检索到的知识库资料，不出现知识库外事实性补充。'],
['来源引用','每个关键结论至少有一个可打开来源；来源片段能支持该结论。'],
['资料不足','无相关或依据不足时明确提示资料不足，不生成无来源的完整结论。'],
['数据边界','待补充、已删除或不可用资料不得被检索或引用。']], [1.45,5.05])

heading(doc,'八、数据指标与迭代规划',1)
heading(doc,'8.1 指标设计目标',2)
add_para(doc,'V1 数据指标用于评估 Refind 是否帮助用户完成“资料沉淀、知识找回、AI 调用与来源核验”的核心使用链路，并识别资料解析、检索与回答中的主要问题。')
add_code(doc,['创建知识库 → 添加资料 → 搜索或提问 → 获得回答/资料 → 查看来源 → 继续使用'])
heading(doc,'8.2 核心指标',2)
add_table(doc,['模块','指标','定义','观察目的'],[
['知识库创建','知识库创建完成率','成功创建次数 ÷ 点击创建次数','判断创建流程是否顺畅。'],
['资料沉淀','资料保存完成率','成功保存次数 ÷ 发起解析次数','判断解析、预览与保存流程是否顺畅。'],
['资料沉淀','解析成功率','成功获得标题或正文节选资料数 ÷ 发起解析资料数','判断链接解析能力。'],
['AI 结构化','摘要编辑率','被用户修改摘要数 ÷ AI 生成摘要数','判断摘要是否符合用户理解。'],
['AI 结构化','标签编辑率','被新增、删除或修改标签数 ÷ AI 生成标签数','判断标签质量及可用性。'],
['资料找回','搜索结果点击率','点击搜索建议条目次数 ÷ 浮层展示次数','判断检索结果相关性。'],
['AI 调用','AI 对话发起率','至少发起一次对话用户数 ÷ 拥有至少 3 篇可用资料用户数','判断用户是否认可向资料库提问。'],
['AI 调用','有效知识问答完成数','含来源回答后，用户点击来源、继续追问、复制回答或查看原文的次数','衡量回答是否被实际使用。'],
['可信度','引用点击率','点击任一来源引用回答数 ÷ 含引用回答数','判断溯源是否建立信任。'],
['资料不足','资料不足率','返回资料不足回答数 ÷ AI 回答总数','判断资料覆盖度与检索效果。'],
['留存','7 日回访率','首次访问后 7 日内再次访问用户数 ÷ 首次访问用户数','判断持续使用价值。']], [1.05,1.35,2.3,1.8])
heading(doc,'8.3 关键事件埋点',2)
add_table(doc,['事件名称','触发时机','核心属性'],[
['knowledge_base_created','成功创建知识库','知识库名称、是否默认知识库。'],
['material_parse_started','提交链接并开始解析','来源平台、入口位置、目标知识库。'],
['material_parse_completed','链接解析完成','是否提取正文、正文长度、解析耗时。'],
['material_saved','成功保存资料','知识库 ID、来源平台、是否仅保存链接。'],
['search_submitted','输入搜索内容','搜索范围、输入长度。'],
['search_result_clicked','点击搜索建议条目','搜索范围、结果排名、资料 ID。'],
['chat_question_submitted','用户发送问题','是否含标签、标签数量、问题长度。'],
['rag_retrieval_completed','检索完成','召回片段数、涉及知识库数、最高相似度。'],
['chat_answer_completed','AI 回答生成完成','回答耗时、引用数、资料不足标识。'],
['citation_opened','用户打开来源引用','资料 ID、知识库 ID、引用位置。'],
['source_opened','用户查看原文','资料 ID、来源平台。'],
['answer_feedback_submitted','用户提交回答反馈','反馈类型、回答 ID。']], [1.85,1.75,2.9])
heading(doc,'8.4 迭代规划',2)
heading(doc,'V1.0：个人知识库基础闭环',3)
for b in ['多知识库及默认知识库。','链接保存、解析、摘要与标签。','全局及知识库内语义搜索。','基于 RAG 的全局 AI 对话、#标签 聚焦检索、回答引用、来源片段与资料不足提示。']: add_bullet(doc,b)
heading(doc,'V1.5：检索质量与资料管理优化',3)
for b in ['支持资料批量移动、批量删除与批量打标签。','优化网页解析失败率与正文提取质量。','增加混合检索或重排序，提升召回准确性。','支持用户选择本轮对话参考的知识库范围。','支持重新解析、重新生成摘要和标签；优化回答引用颗粒度与片段高亮。','增加回答复制、收藏和历史对话查看。']: add_bullet(doc,b)
heading(doc,'V2.0：个人知识应用升级',3)
for b in ['登录注册、个人数据隔离与跨设备同步。','文件、图片、音频等多模态资料导入。','共享知识库与协作权限。','联网补充模式，清晰区分个人资料与外部来源。','任务式输出，如调研 Brief、学习计划、竞品分析大纲。','在用户授权范围内引入 Agent，用于复杂任务拆解、资料缺口识别与多步骤产出。']: add_bullet(doc,b)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
